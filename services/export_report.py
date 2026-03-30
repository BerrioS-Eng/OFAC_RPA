import os
import logging
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config.settings import settings

logger = logging.getLogger(__name__)

def _ensure_reports_dir():
    """Creates reports directory if it doesn't exist."""
    os.makedirs(settings.reports_dir, exist_ok=True)

def export_report(title, registers):
    """"""
    if not registers:
        logger.info("No data to export.")
        return None

    _ensure_reports_dir()

    wb = Workbook()
    ws = wb.active
    ws.title = title

    # Headers from data keys
    headers = list(registers[0].keys())[:-1]

    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_alignment = Alignment(horizontal="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    ws.freeze_panes = "A2"

    # Data
    data_font = Font(name="Arial", size=10)
    for row, reg in enumerate(registers, 2):
        for col, key in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=reg.get(key, "")).font = data_font
    
    # Auto-filter
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{len(registers) + 1}"

    # Save file
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    normalized_title = title.replace(" ", "_")
    normalized_title = "".join(c for c in normalized_title if c.isalnum() or c == "_")
    filename = f"{normalized_title}_{timestamp}.xlsx"
    filepath = os.path.join(settings.reports_dir, filename)

    wb.save(filepath)

    logger.info(f"Report exported: {filepath} with {len(registers)} registers.")
    return filepath


